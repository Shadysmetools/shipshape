#!/usr/bin/env python3
"""Measure real working hours from Claude Code session logs, map them to tasks,
and write an Excel workbook plus Odoo account.analytic.line import rows.

The method (also written into the workbook's Method sheet):
  1. Read every transcript line with a timestamp from the main session file(s)
     AND their subagent files (<session>/subagents/*.jsonl). By default only
     "user" and "assistant" lines count; "queue-operation" lines (messages
     queued into an idle session) are always excluded because they inflate
     totals with time nobody worked.
  2. SORT each file's timestamps. (Unsorted input once produced negative
     intervals and negative hours; tests/test_timesheet.py guards this.)
  3. Within one file, two consecutive lines at most GAP minutes apart make an
     active interval. Longer gaps are breaks.
  4. Union the active intervals of all files, so parallel subagents count once.
  5. Cut the union into work blocks wherever the break is longer than GAP.
  6. Attach the typed requests (main transcripts) and git commits (+/- 5 min)
     to each block, then map the block to a task with the rules in tasks.json.
  7. Split by local calendar day, round each (day, task) to the grid (0.25 h),
     run the self-checks, and write the workbook. Values only, no formulas.

This is session time, not keyboard time. Never ship estimates as hours.
Nothing here writes to Odoo: the import rows are a file for a human to review.

Human-equivalent hours (shown next to measured hours on every row; booked and
imported only with --mode human):
  H1. Count what was done in each block from the logs: tool calls (scenario runs,
      UI steps, bugs filed), distinct files changed, commits, typed requests
      (meetings). What counts as what is a regex in the rates file.
  H2. Minutes = sum(count x rate per activity), never below the block minimum.
      Identical repeated tool calls count once; a block with nothing counted is 0.
  H3. Share the minutes across the block's tasks like measured time, times the
      task's complexity multiplier; book on the block's start day.
  H4. Every row carries the reasoning (counts x rates). The rates file and the
      team's rule (timesheet.rule) are written into the Method sheet.
"""
from __future__ import annotations

import argparse
import collections
import csv
import datetime as dt
import glob
import json
import os
import re
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from privacy import mask_text  # noqa: E402

COUNTED_TYPES = ("user", "assistant")
NEVER_COUNTED = ("queue-operation",)
FRAGMENT_SEC = 600  # blocks shorter than this inherit the nearest same-day task
DEFAULT_RATES = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates", "human-rates.example.json")
SOURCES = ("tool", "file", "commit", "request")


# ---------------------------------------------------------------- reading
def parse_ts(s: str) -> float:
    return dt.datetime.fromisoformat(s.replace("Z", "+00:00")).timestamp()


def _text_of(msg) -> str:
    c = msg.get("content") if isinstance(msg, dict) else None
    if isinstance(c, str):
        return c
    if isinstance(c, list):
        parts = []
        for p in c:
            if isinstance(p, dict) and p.get("type") == "text":
                parts.append(p.get("text", ""))
            elif isinstance(p, dict) and p.get("type") == "tool_result":
                return ""  # a tool result is not a typed request
        return " ".join(parts)
    return ""


def read_file(path: str, types=COUNTED_TYPES, main: bool = True):
    """Return (sorted timestamps, typed prompts [(t, text)]) for one jsonl file."""
    stamps, prompts = [], []
    with open(path, encoding="utf-8", errors="replace") as fh:
        for line in fh:
            try:
                d = json.loads(line)
            except ValueError:
                continue
            if not isinstance(d, dict):
                continue
            kind, ts = d.get("type"), d.get("timestamp")
            if not ts or kind in NEVER_COUNTED or (types and kind not in types):
                continue
            try:
                t = parse_ts(ts)
            except ValueError:
                continue
            stamps.append(t)
            if main and kind == "user" and not d.get("isMeta"):
                txt = re.sub(r"<[^>]+>[\s\S]*?</[^>]+>", " ", _text_of(d.get("message"))).strip()
                if txt and not txt.startswith("This session is being continued"):
                    prompts.append((t, " ".join(txt.split())[:200]))
    stamps.sort()  # REQUIRED: lines are not guaranteed to be in time order
    prompts.sort()
    return stamps, prompts


def read_actions(path: str):
    """Tool calls in one jsonl file -> sorted [(t, tool name, input as text, file path or None)]."""
    out = []
    with open(path, encoding="utf-8", errors="replace") as fh:
        for line in fh:
            try:
                d = json.loads(line)
            except ValueError:
                continue
            if not isinstance(d, dict) or d.get("type") != "assistant" or not d.get("timestamp"):
                continue
            content = (d.get("message") or {}).get("content") if isinstance(d.get("message"), dict) else None
            if not isinstance(content, list):
                continue
            try:
                t = parse_ts(d["timestamp"])
            except ValueError:
                continue
            for p in content:
                if isinstance(p, dict) and p.get("type") == "tool_use":
                    inp = p.get("input") if isinstance(p.get("input"), dict) else {}
                    fp = inp.get("file_path") or inp.get("notebook_path")
                    out.append((t, str(p.get("name", "")), json.dumps(inp, ensure_ascii=False)[:2000], fp))
    out.sort(key=lambda x: x[0])
    return out


def session_files(project_dir: str, sessions=None):
    """Main + subagent transcript files for the given session ids (or all)."""
    out = []
    mains = sorted(glob.glob(os.path.join(project_dir, "*.jsonl")))
    for m in mains:
        sid = os.path.splitext(os.path.basename(m))[0]
        if sessions and sid not in sessions:
            continue
        out.append((sid, "main", m))
        for s in sorted(glob.glob(os.path.join(project_dir, sid, "subagents", "*.jsonl"))):
            out.append((sid, "subagent", s))
    return out


# ---------------------------------------------------------------- measuring
def file_intervals(stamps, gap_sec):
    """Active intervals inside one file. Input MUST be sorted; we sort defensively."""
    s = sorted(stamps)
    return [(a, b) for a, b in zip(s, s[1:]) if 0 < b - a <= gap_sec]


def union(intervals):
    merged = []
    for a, b in sorted(intervals):
        if merged and a <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], b)
        else:
            merged.append([a, b])
    return [(a, b) for a, b in merged]


def blocks_from(merged, gap_sec):
    blocks = []
    for a, b in merged:
        if blocks and a - blocks[-1]["end"] <= gap_sec:
            blocks[-1]["end"] = b
            blocks[-1]["sec"] += b - a
            blocks[-1]["parts"].append((a, b))
        else:
            blocks.append({"start": a, "end": b, "sec": b - a, "parts": [(a, b)]})
    return blocks


def measure(per_file_stamps, gap_min=20):
    """per_file_stamps: list of timestamp lists (one per transcript file).
    Returns (merged intervals, blocks, total hours)."""
    gap = gap_min * 60
    iv = []
    for stamps in per_file_stamps:
        iv += file_intervals(stamps, gap)
    merged = union(iv)
    blocks = blocks_from(merged, gap)
    # invariants: a violation stops the run instead of producing a wrong timesheet
    assert all(b >= a for a, b in merged), "negative merged interval"
    assert all(x["sec"] >= 0 for x in blocks), "negative block"
    assert all(n["start"] >= p["end"] for p, n in zip(blocks, blocks[1:])), "overlapping blocks after union"
    total = sum(b - a for a, b in merged) / 3600
    return merged, blocks, total


def split_days(a, b, tz):
    """Split [a, b) (epoch seconds) by local calendar day -> [(date_iso, seconds)]."""
    out = []
    while a < b:
        d = dt.datetime.fromtimestamp(a, tz)
        nxt = (d.replace(hour=0, minute=0, second=0, microsecond=0) + dt.timedelta(days=1)).timestamp()
        e = min(b, nxt)
        out.append((d.date().isoformat(), e - a))
        a = e
    return out


def parse_tz(spec: str):
    m = re.fullmatch(r"([+-])(\d{1,2}):?(\d{2})?", spec.strip())
    if m:
        sign = 1 if m.group(1) == "+" else -1
        return dt.timezone(sign * dt.timedelta(hours=int(m.group(2)), minutes=int(m.group(3) or 0)))
    try:
        from zoneinfo import ZoneInfo
        return ZoneInfo(spec)
    except Exception as exc:  # noqa: BLE001
        raise SystemExit(f"unknown time zone {spec!r}: use +03:00 style or an IANA name ({exc})")


# ---------------------------------------------------------------- git
def git_commits(repos, author=None):
    commits = []
    for repo in repos:
        cmd = ["git", "-C", repo, "log", "--all", "--format=%H|%aI|%s"]
        if author:
            cmd.append(f"--author={author}")
        out = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120).stdout
        for line in out.splitlines():
            try:
                h, d, sub = line.split("|", 2)
                commits.append((parse_ts(d), h[:7], sub[:120], os.path.basename(os.path.normpath(repo))))
            except ValueError:
                continue
    return sorted(set(commits))


# ---------------------------------------------------------------- mapping
def _first_task(text, tasks):
    for t in tasks:
        for rx in t.get("match", []):
            if re.search(rx, text, re.I):
                return t["id"]
    return None


def assign(blocks, cfg, tz):
    """Give every block an allocation {task_id: seconds} and a basis string."""
    tasks = cfg.get("tasks", [])
    excl = [re.compile(r, re.I) for r in cfg.get("exclude", [])]
    manual = cfg.get("manual", {})
    for b in blocks:
        key = dt.datetime.fromtimestamp(b["start"], tz).strftime("%Y-%m-%d %H:%M")
        b["key"] = key
        if key in manual:
            tid = manual[key]
            b["alloc"] = {} if tid is None else {tid: b["sec"]}
            b["basis"] = "manual map" if tid is not None else "manual map: excluded"
            continue
        prompts_text = b.get("prompts_text", [])
        texts = prompts_text + b.get("commit_subjects", [])
        # every typed request is admin (e.g. "build my timesheet") and nothing was committed -> excluded
        if prompts_text and not b.get("commit_subjects") and all(any(r.search(x) for r in excl) for x in prompts_text):
            b["alloc"], b["basis"] = {}, "excluded: admin / non-project (exclude rule)"
            continue
        votes = collections.Counter()
        for x in texts:
            if any(r.search(x) for r in excl):
                continue
            tid = _first_task(x, tasks)
            if tid is not None:
                votes[tid] += 1
        if votes:
            n = sum(votes.values())
            b["alloc"] = {t: b["sec"] * k / n for t, k in votes.items()}
            b["basis"] = f"{len(b.get('commit_subjects', []))} commits + {len(b.get('prompts_text', []))} requests matched"
        else:
            b["alloc"], b["basis"] = None, "unassigned"
    # fragments under FRAGMENT_SEC inherit the nearest assigned block of the same day
    for i, b in enumerate(blocks):
        if b["alloc"] is not None:
            continue
        day = b["key"][:10]
        cands = [x for x in blocks if x is not b and x["key"][:10] == day and x["alloc"]]
        if b["sec"] < FRAGMENT_SEC and cands:
            near = min(cands, key=lambda x: abs(x["start"] - b["start"]))
            tot = sum(near["alloc"].values())
            b["alloc"] = {t: b["sec"] * s / tot for t, s in near["alloc"].items()}
            b["basis"] = f"fragment < {FRAGMENT_SEC // 60} min: inherits block {near['key']}"
        elif cfg.get("default_task") is not None:
            b["alloc"] = {cfg["default_task"]: b["sec"]}
            b["basis"] = "default task (no rule matched)"
    return blocks


def round_to(h, grid):
    return round(round(h / grid) * grid, 2) if grid else round(h, 2)


# ---------------------------------------------------------------- human-equivalent
def load_rates(path):
    """Read and validate a rates file; a bad file stops the run with one line."""
    try:
        rates = json.load(open(path, encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise SystemExit(f"rates error in {path}: {exc}")
    acts = rates.get("activities") if isinstance(rates, dict) else None
    if not isinstance(acts, dict) or not acts:
        raise SystemExit(f"rates error in {path}: 'activities' must be a non-empty mapping")
    for name, a in acts.items():
        if not isinstance(a, dict):
            raise SystemExit(f"rates error at activities.{name}: expected a mapping")
        if not isinstance(a.get("minutes"), (int, float)) or a["minutes"] < 0:
            raise SystemExit(f"rates error at activities.{name}.minutes: expected a number >= 0, got {a.get('minutes')!r}")
        if a.get("source") not in SOURCES:
            raise SystemExit(f"rates error at activities.{name}.source: expected one of {', '.join(SOURCES)}, got {a.get('source')!r}")
        for rx in [a.get("tool", "")] + list(a.get("match", [])):
            try:
                re.compile(rx)
            except re.error as exc:
                raise SystemExit(f"rates error at activities.{name}: bad regex {rx!r} ({exc})")
    m = rates.get("block_minimum_minutes", 0)
    if not isinstance(m, (int, float)) or m < 0:
        raise SystemExit(f"rates error at block_minimum_minutes: expected a number >= 0, got {m!r}")
    return rates


def task_multiplier(rates, tid):
    v = (rates.get("task_multiplier") or {}).get(str(tid), 1)
    return float(v) if isinstance(v, (int, float)) and v > 0 else 1.0


def count_activities(b, rates):
    """{activity: count} for one block, from its tool calls, commits and typed requests."""
    counts = {}
    for name, a in rates["activities"].items():
        src, tool_rx, rxs = a["source"], a.get("tool") or ".", [re.compile(r, re.I) for r in a.get("match", [])]
        if src == "commit":
            n = len(b.get("commit_subjects", []))
        elif src == "request":
            n = sum(1 for x in b.get("prompts_text", []) if any(r.search(x) for r in rxs))
        else:
            hits = [x for x in b.get("actions", []) if re.search(tool_rx, x[1], re.I) and (not rxs or any(r.search(x[2]) for r in rxs))]
            if src == "file":
                n = len({x[3] for x in hits if x[3]})
            else:  # the agent re-running the same command is iteration, not new work: identical calls count once
                n = len(hits) if a.get("distinct") is False else len({(x[1], x[2]) for x in hits})
        if n:
            counts[name] = n
    return counts


def human_block(b, rates):
    """(minutes, reasoning) a human would spend on one block, before task shares."""
    counts = count_activities(b, rates)
    acts = rates["activities"]
    parts = [f"{n} {acts[k].get('label', k)} x {acts[k]['minutes']:g}" for k, n in counts.items()]
    raw = sum(n * acts[k]["minutes"] for k, n in counts.items())
    if not counts:
        return 0, "no counted activity: 0 min"
    floor = rates.get("block_minimum_minutes", 0)
    minutes = max(raw, floor)
    why = " + ".join(parts) + f" = {raw:g} min"
    if minutes > raw:
        why += f" -> block minimum {floor:g} min"
    return minutes, why


# ---------------------------------------------------------------- checks
def run_checks(daily, booked, measured_total, cfg, existing=None, lookup=None, max_day=16.0, grid=0.25, unassigned=0, basis="measured"):
    checks = []

    def add(name, ok, detail):
        checks.append({"check": name, "result": "PASS" if ok is True else ("FAIL" if ok is False else "NOT CHECKED"), "detail": detail})

    booked_total = round(sum(booked.values()), 2)
    tol = grid / 2 * max(1, len(booked)) if grid else 0.01
    add("Totals match", abs(booked_total - measured_total) <= tol + 1e-9,
        f"booked {booked_total:.2f} h vs {basis} {measured_total:.2f} h (rounding tolerance {tol:.2f} h)")
    neg = [k for k, v in booked.items() if v < 0] + [k for k, v in daily.items() if v < 0]
    add("No negative hours", not neg, "none" if not neg else f"{len(neg)} negative rows")
    per_day = collections.defaultdict(float)
    for (d, _t), h in booked.items():
        per_day[d] += h
    ex_day = collections.defaultdict(float)
    for r in existing or []:
        ex_day[r["date"]] += float(r.get("hours", 0))
    over = sorted(d for d in set(per_day) | set(ex_day) if per_day[d] + ex_day[d] > max_day)
    add(f"No day over {max_day:g} h (new + existing)", not over, "none" if not over else ", ".join(over))
    if existing is None:
        add("No overlap with existing lines", None, "no existing-lines export given (--existing)")
    else:
        clash = sorted({(r["date"], r.get("task_id")) for r in existing} & set(booked))
        add("No overlap with existing lines", not clash,
            "none" if not clash else "; ".join(f"{d} task {t}" for d, t in clash))
    emp = cfg.get("employee", {})
    if lookup is None:
        add("Employee exists and is visible", None, "no lookup export given (--lookup); importer must confirm")
    else:
        hits = [e for e in lookup.get("employees", [])
                if (emp.get("id") and e.get("id") == emp.get("id")) or (emp.get("name") and e.get("name") == emp.get("name"))]
        comp = cfg.get("company_id")
        ok = bool(hits) and (comp is None or any(e.get("company_id") == comp for e in hits))
        add("Employee exists and is visible", ok,
            f"{len(hits)} match(es) among {len(lookup.get('employees', []))} visible employees" + (f", company {comp}" if comp else ""))
    add("Every block assigned", unassigned == 0, "none" if not unassigned else f"{unassigned} block(s) unassigned: add a rule or a manual entry")
    return checks


def summary_line(checks):
    n = sum(c["result"] == "PASS" for c in checks)
    m = sum(c["result"] == "FAIL" for c in checks)
    k = sum(c["result"] == "NOT CHECKED" for c in checks)
    return f"{n} pass / {m} fail / {k} not checked"


# ---------------------------------------------------------------- output
METHOD = __doc__.strip().splitlines()


def rates_lines(rates, path):
    out = [f"Rates file: {path}", f"Block minimum: {rates.get('block_minimum_minutes', 0):g} min"]
    for name, a in rates["activities"].items():
        how = f"source {a['source']}" + (f", tool /{a['tool']}/" if a.get("tool") else "") + (f", match {a['match']}" if a.get("match") else "")
        out.append(f"  {name} ({a.get('label', name)}): {a['minutes']:g} min each; {how}")
    mult = {k: v for k, v in (rates.get("task_multiplier") or {}).items() if isinstance(v, (int, float))}
    out.append(f"Task complexity multipliers: {mult or 'none'}")
    return out


def write_workbook(out, ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    bold = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="DDDDDD")

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, hdr_fill
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(80, max(10, width + 2))
        return ws

    names = ctx["names"]
    per_task = collections.defaultdict(float)
    for (d, t), h in ctx["booked"].items():
        per_task[t] += h
    tids = sorted(set(per_task) | set(ctx["measured_task"]) | set(ctx["human_task"]), key=str)
    srows = [(t, names.get(t, ""), round(per_task.get(t, 0), 2), round(ctx["measured_task"].get(t, 0), 2), round(ctx["human_task"].get(t, 0), 2))
             for t in tids]
    srows.append(("TOTAL", "", round(sum(per_task.values()), 2), round(ctx["measured_total"], 2), round(ctx["human_total"], 2)))
    srows.append(("Excluded (admin / non-project)", "", "", round(ctx["excluded"], 2), ""))
    srows.append(("Checks", ctx["check_line"], "", "", ""))
    srows.append(("Mode", f"{ctx['mode']}: booked hours = {ctx['basis']} hours", "", "", ""))
    srows.append(("Rule", ctx["rule"] or "(none given: set timesheet.rule)", "", "", ""))
    srows.append(("Gap rule", f"{ctx['gap']} min", "", "", ""))
    srows.append(("Time zone", ctx["tz_name"], "", "", ""))
    sheet("Summary", ["Task ID", "Task", "Booked hours", "Measured hours", "Human-equivalent hours"], srows, first=True)
    keys = sorted(set(ctx["daily"]) | set(ctx["human_daily"]) | set(ctx["booked"]), key=lambda x: (x[0], str(x[1])))
    sheet("Daily", ["Date", "Task ID", "Task", "Measured hours", "Human-equivalent hours", "Booked hours", "Human-equivalent reasoning"],
          [(d, t, names.get(t, ""), round(ctx["daily"].get((d, t), 0), 2), round(ctx["human_daily"].get((d, t), 0), 2),
            ctx["booked"].get((d, t), 0), "; ".join(ctx["human_why"].get((d, t), []))) for d, t in keys])
    sheet("Sessions", ["Session", "Kind", "File", "Lines counted", "First (local)", "Last (local)"], ctx["sessions"])
    sheet("Evidence", ["Block start", "End", "Active hours", "Human-equivalent hours", "Allocation", "Basis", "Human-equivalent reasoning",
                       "Requests (masked)", "Commits (masked)"], ctx["evidence"])
    sheet("Method", ["Method"], [(line,) for line in METHOD + ["", f"Gap: {ctx['gap']} min. Time zone: {ctx['tz_name']}. Grid: {ctx['grid']} h. Max per day: {ctx['max_day']} h.",
                                                       f"Mode: {ctx['mode']}. Booked and imported hours are {ctx['basis']} hours.",
                                                       f"Rule (timesheet.rule): {ctx['rule'] or '(none given)'}"] + ctx["rates_lines"] +
                                                      ["Reproduce: python scripts/timesheet.py with the same arguments (see Summary)."]])
    sheet("Checks", ["Check", "Result", "Detail"], [(c["check"], c["result"], c["detail"]) for c in ctx["checks"]])
    sheet("Odoo Import", ["Date", "Project/Database ID", "Project", "Task/Database ID", "Task", "Employee", "Description", "Hours Spent"], ctx["import_rows"])
    wb.save(out)


def load_config(path):
    """The timesheet section of qa.config.yaml: {mode, rule, rates}. Missing file or section -> {}."""
    if not path:
        return {}
    import yaml
    try:
        data = yaml.safe_load(open(path, encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError) as exc:
        raise SystemExit(f"config error in {path}: {exc}")
    sec = data.get("timesheet") if isinstance(data, dict) else None
    if sec is None:
        return {}
    if not isinstance(sec, dict):
        raise SystemExit(f"config error at timesheet: expected a mapping (mode, rule, rates)")
    if sec.get("mode") not in (None, "measured", "human"):
        raise SystemExit(f"config error at timesheet.mode: expected measured or human, got {sec['mode']!r}")
    return sec


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--project-dir", required=True, help="~/.claude/projects/<slug> folder with the transcripts")
    ap.add_argument("--session", action="append", help="session id (repeatable); default: all sessions in the folder")
    ap.add_argument("--tasks", required=True, help="tasks.json: project, employee, tasks with regex rules, exclude, manual")
    ap.add_argument("--repo", action="append", default=[], help="git repo to read commits from (repeatable)")
    ap.add_argument("--author", help="only commits by this git author")
    ap.add_argument("--gap", type=int, choices=(10, 15, 20), default=20, help="minutes; default 20")
    ap.add_argument("--tz", default="+00:00", help="local time zone: +03:00 or an IANA name")
    ap.add_argument("--since", help="YYYY-MM-DD (local, inclusive)")
    ap.add_argument("--until", help="YYYY-MM-DD (local, inclusive)")
    ap.add_argument("--grid", type=float, default=0.25, help="rounding grid in hours (0 = none)")
    ap.add_argument("--max-day", type=float, default=16.0)
    ap.add_argument("--existing", help="JSON export of existing analytic lines [{date, task_id, hours}] (read-only export)")
    ap.add_argument("--lookup", help='JSON {"employees": [{id, name, company_id}]} visible to the importing user')
    ap.add_argument("--mode", choices=("measured", "human"), help="hours to book and import: measured (default) or human-equivalent")
    ap.add_argument("--rates", help="human-equivalent rates JSON; default templates/human-rates.example.json")
    ap.add_argument("--config", help="qa.config.yaml: reads timesheet.mode, timesheet.rule, timesheet.rates")
    ap.add_argument("--out", required=True, help="output .xlsx; .csv and .json import files are written next to it")
    a = ap.parse_args(argv)

    tcfg = load_config(a.config)
    mode = a.mode or tcfg.get("mode") or "measured"
    rates_path = a.rates or tcfg.get("rates") or DEFAULT_RATES
    rates = load_rates(rates_path)
    rule = str(tcfg.get("rule") or rates.get("rule") or "").strip()
    basis = "human-equivalent" if mode == "human" else "measured"

    tz = parse_tz(a.tz)
    cfg = json.load(open(a.tasks, encoding="utf-8"))
    names = {t["id"]: t.get("name", "") for t in cfg.get("tasks", [])}
    files = session_files(os.path.expanduser(a.project_dir), set(a.session) if a.session else None)
    if not files:
        raise SystemExit("no transcripts found")
    lo = dt.datetime.fromisoformat(a.since).replace(tzinfo=tz).timestamp() if a.since else float("-inf")
    hi = (dt.datetime.fromisoformat(a.until).replace(tzinfo=tz) + dt.timedelta(days=1)).timestamp() if a.until else float("inf")

    per_file, prompts, sessions, actions = [], [], [], []
    for sid, kind, path in files:
        st, pr = read_file(path, main=(kind == "main"))
        st = [t for t in st if lo <= t < hi]
        per_file.append(st)
        prompts += [p for p in pr if lo <= p[0] < hi]
        actions += [x for x in read_actions(path) if lo <= x[0] < hi]
        if st:
            f = lambda t: dt.datetime.fromtimestamp(t, tz).strftime("%Y-%m-%d %H:%M")  # noqa: E731
            sessions.append((sid[:8], kind, os.path.basename(path), len(st), f(st[0]), f(st[-1])))
    merged, blocks, measured_total = measure(per_file, a.gap)
    commits = git_commits(a.repo, a.author) if a.repo else []
    for b in blocks:
        b["prompts_text"] = [t for ts, t in prompts if b["start"] <= ts <= b["end"] + 1]
        b["commit_subjects"] = [f"{c[2]}" for c in commits if b["start"] - 300 <= c[0] <= b["end"] + 300]
        b["commit_refs"] = [f"{c[3]}@{c[1]} {c[2]}" for c in commits if b["start"] - 300 <= c[0] <= b["end"] + 300]
        b["actions"] = [x for x in actions if b["start"] <= x[0] <= b["end"] + 1]
    assign(blocks, cfg, tz)

    daily = collections.defaultdict(float)
    measured_task = collections.defaultdict(float)
    human_daily = collections.defaultdict(float)
    human_task = collections.defaultdict(float)
    human_why = collections.defaultdict(list)
    excluded = 0.0
    unassigned = 0
    evidence = []
    for b in blocks:
        if b["alloc"] is None:
            unassigned += 1
        elif not b["alloc"]:
            excluded += b["sec"] / 3600
        share = {t: s / b["sec"] for t, s in (b["alloc"] or {}).items()} if b["sec"] else {}
        for a0, a1 in b["parts"]:
            for d, sec in split_days(a0, a1, tz):
                for t, f in share.items():
                    daily[(d, t)] += sec * f / 3600
                    measured_task[t] += sec * f / 3600
        # human-equivalent: counted activities x rates, shared like measured time, booked on the block's start day
        minutes, why = human_block(b, rates)
        block_h = 0.0
        for t, f in share.items():
            m = task_multiplier(rates, t)
            h = minutes / 60 * f * m
            block_h += h
            human_daily[(b["key"][:10], t)] += h
            human_task[t] += h
            human_why[(b["key"][:10], t)].append(f"{b['key'][11:]} {why}" + (f" x share {f:.2f}" if f < 0.999 else "")
                                                 + (f" x complexity {m:g}" if m != 1 else "") + f" = {h:.2f} h")
        if b["alloc"] is None:
            block_h = minutes / 60  # shown for review; an unassigned block is never booked
        evidence.append((b["key"], dt.datetime.fromtimestamp(b["end"], tz).strftime("%H:%M"), round(b["sec"] / 3600, 2), round(block_h, 2),
                         "; ".join(f"{t}: {s / 3600:.2f} h" for t, s in (b["alloc"] or {}).items()) or ("UNASSIGNED" if b["alloc"] is None else "excluded"),
                         b["basis"], "excluded: not billed" if b["alloc"] == {} else why,
                         mask_text(" | ".join(b["prompts_text"])[:500]), mask_text(" | ".join(b["commit_refs"])[:500])))
    measured_billable = sum(daily.values())
    human_total = sum(human_daily.values())
    source = human_daily if mode == "human" else daily
    billable = human_total if mode == "human" else measured_billable
    booked = {k: round_to(v, a.grid) for k, v in source.items()}
    booked = {k: v for k, v in booked.items() if v > 0}
    existing = json.load(open(a.existing, encoding="utf-8")) if a.existing else None
    lookup = json.load(open(a.lookup, encoding="utf-8")) if a.lookup else None
    checks = run_checks(source, booked, billable, cfg, existing, lookup, a.max_day, a.grid, unassigned, basis)
    if mode == "human":
        bare = [k for k in booked if not human_why.get(k)]
        checks.append({"check": "Every human-equivalent row has its reasoning", "result": "PASS" if not bare else "FAIL",
                       "detail": "none missing" if not bare else f"{len(bare)} row(s) without reasoning"})
    proj = cfg.get("project", {})
    emp = cfg.get("employee", {})
    desc = {t["id"]: t.get("description", t.get("name", "")) for t in cfg.get("tasks", [])}
    import_rows = [(d, proj.get("id"), proj.get("name"), t, names.get(t, ""), emp.get("name"), desc.get(t, ""), h)
                   for (d, t), h in sorted(booked.items(), key=lambda x: (x[0][0], str(x[0][1])))]
    ctx = dict(names=names, booked=booked, daily=daily, measured_total=measured_billable, measured_task=measured_task, excluded=excluded,
               human_daily=human_daily, human_task=human_task, human_total=human_total, human_why=human_why,
               mode=mode, basis=basis, rule=rule, rates_lines=rates_lines(rates, rates_path),
               sessions=sessions, evidence=evidence, checks=checks, check_line=summary_line(checks), import_rows=import_rows,
               gap=a.gap, tz_name=a.tz, grid=a.grid, max_day=a.max_day)
    write_workbook(a.out, ctx)
    base = os.path.splitext(a.out)[0]
    with open(base + "-odoo-import.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Date", "Project/Database ID", "Project", "Task/Database ID", "Task", "Employee", "Description", "Hours Spent"])
        w.writerows(import_rows)
    lines = [{"date": d, "project_id": proj.get("id"), "task_id": t, "employee_id": emp.get("id"), "name": desc.get(t, ""), "unit_amount": h}
             for (d, t), h in sorted(booked.items(), key=lambda x: (x[0][0], str(x[0][1])))]
    with open(base + "-analytic-lines.json", "w", encoding="utf-8") as fh:
        json.dump({"model": "account.analytic.line", "write_status": "NOT WRITTEN: needs explicit owner OK",
                   "hours_basis": basis, "rule": rule, "rows": lines}, fh, indent=1)
    print(f"measured (all) {measured_total:.2f} h | billable {measured_billable:.2f} h | human-equivalent {human_total:.2f} h | "
          f"booked {sum(booked.values()):.2f} h ({basis}) | excluded {excluded:.2f} h | blocks {len(blocks)} | unassigned {unassigned} | gap {a.gap} min")
    print("checks:", summary_line(checks))
    for c in checks:
        if c["result"] != "PASS":
            print(f"  {c['result']}: {c['check']}: {c['detail']}")
    print("wrote", a.out)
    return 0 if all(c["result"] != "FAIL" for c in checks) else 1


if __name__ == "__main__":
    sys.exit(main())
