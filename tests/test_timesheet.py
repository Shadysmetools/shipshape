"""Timesheet measurement. The first test is the regression test for the bug where
unsorted transcript lines produced negative intervals and negative hours."""
import datetime as dt
import json
import os

import pytest
from openpyxl import load_workbook

import timesheet as ts

T0 = dt.datetime(2026, 1, 5, 9, 0, tzinfo=dt.timezone.utc).timestamp()


def iso(sec):
    return dt.datetime.fromtimestamp(T0 + sec, dt.timezone.utc).isoformat().replace("+00:00", "Z")


def write_jsonl(path, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        for r in rows:
            fh.write(json.dumps(r) + "\n")


def line(sec, kind="assistant", text=None):
    d = {"type": kind, "timestamp": iso(sec)}
    if text is not None:
        d["message"] = {"role": "user", "content": text}
    return d


# ---------------------------------------------------------------- the sorting regression
def test_unsorted_lines_never_give_negative_hours(tmp_path):
    """Lines written out of order (resumed sessions, merged files) used to be paired
    in file order: (10 min, 2 min) gave a -8 min "interval". Sorting fixes it."""
    p = tmp_path / "s.jsonl"
    # true timeline: 0, 2, 4, 6, 8, 10 min -> 10 active minutes; written shuffled
    write_jsonl(str(p), [line(m * 60) for m in (10, 2, 0, 8, 4, 6)])
    stamps, _ = ts.read_file(str(p))
    assert stamps == sorted(stamps)
    iv = ts.file_intervals(stamps, 20 * 60)
    assert all(b > a for a, b in iv)
    _, _, total = ts.measure([stamps], 20)
    assert total == pytest.approx(10 / 60)


def test_file_intervals_sorts_defensively():
    raw = [T0 + 600, T0, T0 + 300]  # unsorted input straight into the interval builder
    iv = ts.file_intervals(raw, 1200)
    assert iv == [(T0, T0 + 300), (T0 + 300, T0 + 600)]


def test_interleaved_main_and_subagent_count_parallel_time_once():
    main = [T0 + m * 60 for m in range(0, 31, 5)]       # 09:00-09:30, every 5 min
    sub = [T0 + m * 60 for m in range(10, 41, 2)]       # 09:10-09:40 in parallel
    _, blocks, total = ts.measure([main, sub], 20)
    assert total == pytest.approx(40 / 60)               # union 09:00-09:40, not 30 + 30
    assert len(blocks) == 1


def test_queue_operation_lines_are_excluded(tmp_path):
    p = tmp_path / "q.jsonl"
    write_jsonl(str(p), [line(0), line(300)] + [line(600 + i * 300, "queue-operation") for i in range(12)])
    stamps, _ = ts.read_file(str(p))
    assert len(stamps) == 2
    assert ts.measure([stamps], 20)[2] == pytest.approx(5 / 60)


@pytest.mark.parametrize("gap,expected_min", [(10, 10), (15, 22), (20, 22)])
def test_gap_rule_is_configurable(gap, expected_min):
    stamps = [T0, T0 + 600, T0 + 600 + 12 * 60]  # a 10-min step then a 12-min step
    assert ts.measure([stamps], gap)[2] == pytest.approx(expected_min / 60)


def test_split_days_at_local_midnight():
    tz = dt.timezone(dt.timedelta(hours=3))
    start = dt.datetime(2026, 1, 5, 23, 30, tzinfo=tz).timestamp()
    parts = ts.split_days(start, start + 3600, tz)
    assert parts == [("2026-01-05", 1800), ("2026-01-06", 1800)]


def test_checks_flag_long_day_overlap_and_missing_employee():
    booked = {("2026-01-05", 1): 10.0, ("2026-01-06", 1): 2.0}
    existing = [{"date": "2026-01-05", "task_id": 1, "hours": 7.0}]
    lookup = {"employees": [{"id": 9, "name": "Someone Else", "company_id": 1}]}
    cfg = {"employee": {"id": 7, "name": "Test Developer"}}
    checks = ts.run_checks(booked, booked, 12.0, cfg, existing, lookup)
    by = {c["check"].split(" (")[0]: c["result"] for c in checks}
    assert by["No day over 16 h"] == "FAIL"
    assert by["No overlap with existing lines"] == "FAIL"
    assert by["Employee exists and is visible"] == "FAIL"
    assert by["Totals match"] == "PASS"
    assert ts.summary_line(checks).startswith("3 pass / 3 fail")


def test_checks_not_checked_without_exports():
    checks = ts.run_checks({}, {("d", 1): 1.0}, 1.0, {}, None, None)
    assert sum(c["result"] == "NOT CHECKED" for c in checks) == 2


def test_end_to_end_workbook(tmp_path):
    proj = tmp_path / "proj"
    sid = "11111111-aaaa-bbbb-cccc-000000000001"
    # main: a login-page block 09:00-09:50 and a timesheet-admin block 13:00-13:20
    main = [line(0, "user", "build the login page")] + [line(m * 60) for m in range(5, 51, 5)]
    main += [line(4 * 3600, "user", "prepare my timesheet")] + [line(4 * 3600 + m * 60) for m in range(5, 21, 5)]
    write_jsonl(str(proj / f"{sid}.jsonl"), main[::-1])  # reversed on purpose
    write_jsonl(str(proj / sid / "subagents" / "agent-1.jsonl"), [line(m * 60) for m in range(30, 71, 4)])
    tasks = {"project": {"id": 100, "name": "Demo project"}, "employee": {"id": 7, "name": "Test Developer"},
             "tasks": [{"id": 1, "name": "Login page", "match": ["login"]}], "exclude": ["timesheet"]}
    tf = tmp_path / "tasks.json"
    tf.write_text(json.dumps(tasks))
    lk = tmp_path / "lookup.json"
    lk.write_text(json.dumps({"employees": [{"id": 7, "name": "Test Developer"}]}))
    ex = tmp_path / "existing.json"
    ex.write_text("[]")
    out = tmp_path / "ts.xlsx"
    rc = ts.main(["--project-dir", str(proj), "--tasks", str(tf), "--out", str(out), "--tz", "+00:00",
                  "--gap", "20", "--lookup", str(lk), "--existing", str(ex)])
    assert rc == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Daily", "Sessions", "Evidence", "Method", "Checks", "Odoo Import"]
    rows = list(wb["Odoo Import"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1 and rows[0][3] == 1 and rows[0][7] == pytest.approx(1.25)  # 70 active min -> 1.25 h grid
    for ws in wb.worksheets:  # values, not formulas (Protected View blanks formulas)
        for r in ws.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and v.startswith("=") for v in r)
    checks = {r[0]: r[1] for r in wb["Checks"].iter_rows(min_row=2, values_only=True)}
    assert all(v == "PASS" for v in checks.values()), checks
    lines = json.loads((tmp_path / "ts-analytic-lines.json").read_text())
    assert lines["write_status"].startswith("NOT WRITTEN")
    assert lines["rows"][0]["unit_amount"] == pytest.approx(1.25)


# ---------------------------------------------------------------- human-equivalent mode (issue #5)
RATES = {
    "rule": "Bill what an Odoo tester would spend.",
    "block_minimum_minutes": 15,
    "activities": {
        "scenario": {"label": "scenario run", "minutes": 10, "source": "tool", "tool": "^Bash$", "match": ["record\\.py|pytest"]},
        "step": {"label": "UI step", "minutes": 1, "source": "tool", "tool": "browser_(click|type|navigate)"},
        "bug": {"label": "bug filed", "minutes": 15, "source": "tool", "tool": ".", "match": ["gh issue create"]},
        "file": {"label": "file changed", "minutes": 5, "source": "file", "tool": "^(Edit|Write)$"},
        "commit": {"label": "commit", "minutes": 5, "source": "commit"},
        "meeting": {"label": "meeting", "minutes": 30, "source": "request", "match": ["meeting"]},
    },
    "task_multiplier": {"2": 2},
}


def tool(sec, name, **inp):
    return {"type": "assistant", "timestamp": iso(sec), "message": {"role": "assistant", "content": [
        {"type": "text", "text": "working"}, {"type": "tool_use", "id": f"t{sec}", "name": name, "input": inp}]}}


def test_default_rates_template_is_valid():
    rates = ts.load_rates(ts.DEFAULT_RATES)
    assert rates["rule"] and {"scenario", "step", "bug", "file", "commit", "meeting"} <= set(rates["activities"])


def test_human_block_counts_every_activity(tmp_path):
    p = tmp_path / "a.jsonl"
    rows = [tool(0, "Bash", command="python scripts/record.py plans/login.json"), tool(60, "Bash", command="pytest -q"),
            tool(120, "Bash", command="ls")]                                      # ls is not a scenario
    rows += [tool(180 + i * 10, "mcp__playwright__browser_click", element=f"button {i}") for i in range(5)]
    rows += [tool(240 + i, "Bash", command="pytest -q") for i in range(6)]       # the agent re-running: counts once
    rows += [tool(300, "Bash", command="gh issue create -t \"Save button does nothing\"")]
    rows += [tool(400, "Edit", file_path="a.py"), tool(410, "Edit", file_path="a.py"), tool(420, "Write", file_path="b.py")]
    write_jsonl(str(p), rows[::-1])
    acts = ts.read_actions(str(p))
    assert [x[0] for x in acts] == sorted(x[0] for x in acts)
    b = {"actions": acts, "commit_subjects": ["fix save"], "prompts_text": ["join the meeting notes", "run it"]}
    assert ts.count_activities(b, RATES) == {"scenario": 2, "step": 5, "bug": 1, "file": 2, "commit": 1, "meeting": 1}
    minutes, why = ts.human_block(b, RATES)
    assert minutes == 2 * 10 + 5 * 1 + 15 + 2 * 5 + 5 + 30
    assert "2 scenario run x 10" in why and why.endswith("= 85 min")


def test_human_block_minimum_applies_only_when_something_was_done():
    minutes, why = ts.human_block({"actions": [], "commit_subjects": ["one commit"], "prompts_text": []}, RATES)
    assert minutes == 15 and why == "1 commit x 5 = 5 min -> block minimum 15 min"
    assert ts.human_block({"actions": [], "commit_subjects": [], "prompts_text": ["hi"]}, RATES) == (0, "no counted activity: 0 min")


def test_distinct_false_counts_every_repeat():
    rates = {"activities": {"run": {"minutes": 2, "source": "tool", "tool": "Bash", "distinct": False}}}
    b = {"actions": [(0, "Bash", '{"command": "pytest"}', None)] * 3}
    assert ts.count_activities(b, rates) == {"run": 3}
    rates["activities"]["run"].pop("distinct")
    assert ts.count_activities(b, rates) == {"run": 1}


@pytest.mark.parametrize("bad,msg", [({"activities": {}}, "non-empty"),
                                     ({"activities": {"x": {"minutes": -1, "source": "tool"}}}, "activities.x.minutes"),
                                     ({"activities": {"x": {"minutes": 1, "source": "guess"}}}, "activities.x.source"),
                                     ({"activities": {"x": {"minutes": 1, "source": "tool", "tool": "("}}}, "bad regex")])
def test_bad_rates_file_stops_with_one_line(tmp_path, bad, msg):
    p = tmp_path / "r.json"
    p.write_text(json.dumps(bad))
    with pytest.raises(SystemExit) as e:
        ts.load_rates(str(p))
    assert msg in str(e.value)


def _human_fixture(tmp_path):
    """Two tasks. Login: a 50-min agent block with 1 scenario + 3 UI steps + 1 bug.
    Checkout (complexity 2): a 30-min block that changed 2 files. Then a timesheet-admin block."""
    proj = tmp_path / "proj"
    sid = "22222222-aaaa-bbbb-cccc-000000000002"
    main = [line(0, "user", "test the login page")] + [line(m * 60) for m in range(5, 51, 5)]
    main += [tool(600, "Bash", command="python scripts/record.py plans/login.json")]
    main += [tool(700 + i, "mcp__playwright__browser_type", text=f"field {i}") for i in range(3)]
    main += [tool(900, "Bash", command="gh issue create -t \"login error\"")]
    t2 = 2 * 3600
    main += [line(t2, "user", "fix the checkout total")] + [line(t2 + m * 60) for m in range(5, 31, 5)]
    main += [tool(t2 + 300, "Edit", file_path="cart.py"), tool(t2 + 400, "Edit", file_path="total.py")]
    t3 = 4 * 3600
    main += [line(t3, "user", "prepare my timesheet")] + [line(t3 + m * 60) for m in range(5, 21, 5)]
    main += [tool(t3 + 300, "Bash", command="pytest -q")]                 # excluded block: never billed
    write_jsonl(str(proj / f"{sid}.jsonl"), main[::-1])
    tasks = {"project": {"id": 100, "name": "Demo project"}, "employee": {"id": 7, "name": "Test Developer"},
             "tasks": [{"id": 1, "name": "Login page", "match": ["login"]}, {"id": 2, "name": "Checkout", "match": ["checkout"]}],
             "exclude": ["timesheet"]}
    tf = tmp_path / "tasks.json"
    tf.write_text(json.dumps(tasks))
    rf = tmp_path / "rates.json"
    rf.write_text(json.dumps(RATES))
    cf = tmp_path / "qa.config.yaml"
    cf.write_text("timesheet:\n  mode: human\n  rule: Team rule - bill tester time, not agent time\n")
    lk = tmp_path / "lookup.json"
    lk.write_text(json.dumps({"employees": [{"id": 7, "name": "Test Developer"}]}))
    ex = tmp_path / "existing.json"
    ex.write_text("[]")
    base = ["--project-dir", str(proj), "--tasks", str(tf), "--rates", str(rf), "--tz", "+00:00", "--gap", "20",
            "--lookup", str(lk), "--existing", str(ex)]
    return base, cf


def test_human_mode_books_human_figure_with_reasoning(tmp_path):
    base, cf = _human_fixture(tmp_path)
    out = tmp_path / "human.xlsx"
    assert ts.main(base + ["--config", str(cf), "--out", str(out)]) == 0   # mode comes from timesheet.mode
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Daily", "Sessions", "Evidence", "Method", "Checks", "Odoo Import"]
    # login: 10 + 3 + 15 = 28 min -> 0.5 h; checkout: 2 files x 5 = 10 -> minimum 15 x complexity 2 = 30 min -> 0.5 h
    imp = {r[3]: r[7] for r in wb["Odoo Import"].iter_rows(min_row=2, values_only=True)}
    assert imp == {1: pytest.approx(0.5), 2: pytest.approx(0.5)}
    hdr = [c.value for c in wb["Daily"][1]]
    assert hdr[3:7] == ["Measured hours", "Human-equivalent hours", "Booked hours", "Human-equivalent reasoning"]
    daily = {r[1]: r for r in wb["Daily"].iter_rows(min_row=2, values_only=True)}
    assert daily[1][3] == pytest.approx(50 / 60, abs=0.01) and daily[1][4] == pytest.approx(28 / 60, abs=0.01)
    assert "1 scenario run x 10 + 3 UI step x 1 + 1 bug filed x 15 = 28 min" in daily[1][6]
    assert "block minimum 15 min x complexity 2 = 0.50 h" in daily[2][6]
    checks = {r[0]: r[1] for r in wb["Checks"].iter_rows(min_row=2, values_only=True)}
    assert all(v == "PASS" for v in checks.values()), checks
    assert "Every human-equivalent row has its reasoning" in checks
    method = "\n".join(r[0] or "" for r in wb["Method"].iter_rows(min_row=2, values_only=True))
    assert "Team rule - bill tester time, not agent time" in method and "scenario (scenario run): 10 min each" in method
    summary = {r[0]: r for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["Mode"][1].startswith("human") and summary["Rule"][1].startswith("Team rule")
    assert summary["TOTAL"][2] == pytest.approx(1.0) and summary["TOTAL"][3] == pytest.approx(80 / 60, abs=0.01)
    ev = list(wb["Evidence"].iter_rows(min_row=2, values_only=True))
    assert ev[-1][4] == "excluded" and ev[-1][6] == "excluded: not billed"
    lines = json.loads((tmp_path / "human-analytic-lines.json").read_text())
    assert lines["hours_basis"] == "human-equivalent" and lines["write_status"].startswith("NOT WRITTEN")
    assert sorted(r["unit_amount"] for r in lines["rows"]) == [0.5, 0.5]


def test_measured_mode_imports_measured_but_shows_both(tmp_path):
    base, _ = _human_fixture(tmp_path)
    out = tmp_path / "measured.xlsx"
    assert ts.main(base + ["--out", str(out)]) == 0
    wb = load_workbook(out)
    imp = {r[3]: r[7] for r in wb["Odoo Import"].iter_rows(min_row=2, values_only=True)}
    assert imp == {1: pytest.approx(0.75), 2: pytest.approx(0.5)}            # 50 and 30 measured minutes on the grid
    daily = {r[1]: r for r in wb["Daily"].iter_rows(min_row=2, values_only=True)}
    assert daily[1][4] == pytest.approx(28 / 60, abs=0.01) and daily[1][6]  # human figure + reasoning still shown
    assert json.loads((tmp_path / "measured-analytic-lines.json").read_text())["hours_basis"] == "measured"


def test_cli_mode_overrides_config_and_bad_config_mode_fails(tmp_path):
    base, cf = _human_fixture(tmp_path)
    out = tmp_path / "m.xlsx"
    assert ts.main(base + ["--config", str(cf), "--mode", "measured", "--out", str(out)]) == 0
    assert json.loads((tmp_path / "m-analytic-lines.json").read_text())["hours_basis"] == "measured"
    cf.write_text("timesheet:\n  mode: guess\n")
    with pytest.raises(SystemExit) as e:
        ts.main(base + ["--config", str(cf), "--out", str(out)])
    assert "timesheet.mode" in str(e.value)
